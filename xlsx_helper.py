# -*- coding: utf-8 -*-
"""Generic Excel writer for naks_sm_lib records. Not NAKS-specific beyond
the default column list — pass your own `columns` for a different shape."""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_COLUMNS = [
    ("№ свидетельства", lambda r: r.get("svid_num", "")),
    ("Марка СМ", lambda r: r.get("marka_sm", "")),
    ("Вид СМ", lambda r: r.get("vid_sm", "")),
    ("Производитель", lambda r: r.get("proizvoditel", "")),
    ("Заявитель", lambda r: r.get("zayavitel", "")),
    ("Представитель", lambda r: r.get("predstavitel", "")),
    ("Поставщик", lambda r: r.get("postavshik", "")),
    ("Потребитель (для карточек по партии)", lambda r: r.get("potrebitel", "")),
    ("Диаметр, мм", lambda r: r.get("diametr", "")),
    ("ТУ / ГОСТ", lambda r: r.get("tu_gost", "")),
    ("Способы сварки", lambda r: r.get("sposoby_svarki", "")),
    ("Группы технических устройств", lambda r: r.get("gruppy_tu", "")),
    ("Вид аттестации", lambda r: r.get("vid_attestacii", "")),
    ("№ АЦ", lambda r: r.get("ac_reg_num", "")),
    ("Срок действия до", lambda r: r.get("srok_deystviya", "")),
    ("Аннулировано", lambda r: "да" if r.get("cancelled") else ""),
    ("Примечания (из карточки)", lambda r: r.get("primechaniya", "")),
    ("Ссылка для проверки", lambda r: r.get("_verify_url", "")),
]


def write_xlsx(records, out_path, columns=None, sheet_name="Реестр НАКС",
               summary=None, sort_key=None):
    """
    records: list[dict] — typically parse_detail() outputs, optionally with
             an added "_verify_url" key (see verification_url() in the lib).
    columns: list[(header_title, fn(record)->value)]; defaults to
             DEFAULT_COLUMNS. Build your own subset/order per the user's ask.
    summary: optional list[(label, value)] written to a second "Сводка" sheet.
    sort_key: optional fn(record)->sortable, applied before writing.
    """
    columns = columns or DEFAULT_COLUMNS
    if sort_key:
        records = sorted(records, key=sort_key)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]  # Excel sheet name limit

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col_idx, (title, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    link_col = None
    for col_idx, (title, _) in enumerate(columns, start=1):
        if "ссылка" in title.lower() or "link" in title.lower():
            link_col = col_idx

    for row_idx, rec in enumerate(records, start=2):
        for col_idx, (_, getter) in enumerate(columns, start=1):
            val = getter(rec)
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if col_idx == link_col and val:
                cell.hyperlink = val
                cell.style = "Hyperlink"

    for i in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(records) + 1}"

    if summary:
        ws2 = wb.create_sheet("Сводка")
        ws2["A1"] = "Параметр"
        ws2["B1"] = "Значение"
        ws2["A1"].font = Font(bold=True)
        ws2["B1"].font = Font(bold=True)
        for i, (k, v) in enumerate(summary, start=2):
            ws2.cell(row=i, column=1, value=k)
            c = ws2.cell(row=i, column=2, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws2.column_dimensions["A"].width = 40
        ws2.column_dimensions["B"].width = 70

    wb.save(out_path)
    return out_path
