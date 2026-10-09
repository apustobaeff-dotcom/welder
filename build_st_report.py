# -*- coding: utf-8 -*-
"""
Excel-отчёт по технологиям сварки ВЫСОКОПРОЧНЫХ сталей из реестра НАКС (АЦСТ)
по уже собранным данным (без обращений к сайту):

    python3 build_st_report.py --pool hs_cache.json --details details.json \\
        --out reports/naks_st_vysokoprochnye.xlsx

--pool     кэш collect_high_strength(cache_path=...) — {запрос: {№: строка}}
--details  {№ свидетельства: результат parse_detail()} (можно частичный)

Листы: Сводка · Карточки (с марками СМ, толщинами, подогревом) · Марки СМ
(какие материалы реально стоят в аттестованных технологиях — карта
конкурентов) · Организации (лиды) · Все свидетельства.
"""
import argparse
import collections
import datetime as dt
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, "monitoring"))
import openpyxl  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import naks_st_lib as st  # noqa: E402
from classify import classify  # noqa: E402


def grade_matcher(text):
    c = classify("сталь: " + text)
    return bool(c and (c["max_class"] or 0) >= 420)


H = ["Организация", "СМ: марки", "СМ: как в карточке", "Основные материалы", "Высокопрочные: основание",
     "Способ сварки", "Группы ТУ", "Толщины, мм", "Диаметры, мм", "Подогрев", "Термообработка", "Положения",
     "Технология: шифр, дата", "Технология: название", "№ свидетельства", "Вид аттестации", "Действует до",
     "АЦСТ", "Карточка", "Проверить на naks.ru"]
W = [38, 28, 50, 45, 18, 12, 12, 26, 22, 16, 16, 18, 24, 50, 16, 8, 12, 10, 14, 14]


def _sheet(ws, headers, rows, widths, link_col=None):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3B5C")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    if link_col:
        for r in range(2, ws.max_row + 1):
            c = ws.cell(r, link_col)
            if c.value:
                c.hyperlink, c.value = c.value, "открыть"
                c.font = Font(color="0563C1", underline="single")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pool", required=True)
    ap.add_argument("--details", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()

    raw = json.load(open(a.pool, encoding="utf-8"))
    pool = {k: v for q, rows in raw.items() if q != "_date_from" for k, v in rows.items()}
    hs = {}
    for k, r in pool.items():
        if r["cancelled"]:
            continue
        g, m = st.HIGH_STRENGTH_GROUP in r["groups"], grade_matcher(r["osn_materialy"])
        if g or m:
            r["_hs_reason"] = " + ".join(filter(None, ["группа 3 (М03)" if g else "", "марка σт≥420" if m else ""]))
            hs[k] = r
    det = json.load(open(a.details, encoding="utf-8")) if os.path.exists(a.details) else {}
    for x in det.values():                       # пере-извлечь марки текущей версией разбора
        x["sm_marks"] = st.sm_marks(x.get("sm_text", ""))

    def p(k, f):
        return det.get(k, {}).get("params", {}).get(f, "")

    def row(k):
        r, x = hs[k], det.get(k, {})
        return [r["organization"], ", ".join(x.get("sm_marks", [])), x.get("sm_text", ""),
                x.get("osn_materialy_card") or r["osn_materialy"], r["_hs_reason"], r["sposoby"], r["gruppy"],
                p(k, "Диапазон толщин, мм"), p(k, "Диапазон диаметров, мм"), p(k, "Наличие подогрева"),
                p(k, "Наличие термообработки"), p(k, "Положение при сварке (наплавке)"),
                " ".join(filter(None, [x.get("tech_shifr", ""), x.get("tech_date", "")])), x.get("tech_name", ""),
                k, r["vid_att"], r["srok"], r["ac_name"],
                ("да" if x.get("detail_available") else "недоступна на сайте") if x else "не загружена",
                st.verification_url(k)]

    ready = sorted((k for k in hs if det.get(k, {}).get("detail_available")),
                   key=lambda k: (hs[k]["organization"], k))
    by_org = collections.defaultdict(list)
    for k in hs:
        by_org[hs[k]["organization"]].append(k)
    split = lambda s: {x.strip() for x in s.split(",") if x.strip()}  # noqa: E731
    org_rows = []
    for o, ks in by_org.items():
        marks = collections.Counter(m for k in ks for m in det.get(k, {}).get("sm_marks", []))
        org_rows.append([o, len(ks), sum(1 for k in ks if k in det), ", ".join(m for m, _ in marks.most_common(8)),
                         ", ".join(sorted(set().union(*(split(hs[k]["sposoby"]) for k in ks)))),
                         ", ".join(sorted(set().union(*(split(hs[k]["gruppy"]) for k in ks)))),
                         "; ".join(sorted({hs[k]["osn_materialy"][:60] for k in ks}))[:300]])
    org_rows.sort(key=lambda r: (-r[2], -r[1]))
    mk = collections.defaultdict(lambda: [0, set(), set()])
    for k in ready:
        for m in det[k]["sm_marks"]:
            mk[m][0] += 1
            mk[m][1].add(hs[k]["organization"])
            mk[m][2].update(split(hs[k]["gruppy"]))
    mk_rows = sorted(([m, v[0], len(v[1]), ", ".join(sorted(v[2])), "; ".join(sorted(v[1]))[:600]]
                      for m, v in mk.items()), key=lambda r: -r[1])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сводка"
    nd = sum(1 for k in hs if k in det and not det[k].get("detail_available"))
    for lab, val in [
        ("Реестр", "НАКС: организации, прошедшие проверку готовности к применению аттестованных "
                   "технологий сварки (naks.ru/registry/reg/st/)"),
        ("Дата отчёта", dt.date.today().strftime("%d.%m.%Y")),
        ("Критерий", "действующие, не аннулированные; группа 3 ОМ (М03) или марка с σт ≥ 420 МПа"),
        ("Свидетельств по ВП сталям", len(hs)), ("Организаций", len(by_org)),
        ("Карточек загружено", len(det)), ("  из них с данными", len(ready)), ("  недоступна на сайте", nd),
        ("Лист «Карточки»", "свидетельства с загруженной карточкой: марки СМ, толщины, подогрев, технология"),
        ("Лист «Марки СМ»", "какие СМ реально стоят в аттестованных технологиях ВП сталей — карта конкурентов"),
        ("Лист «Организации»", "лиды: свидетельства, марки СМ, способы, группы ТУ по каждой организации"),
        ("Лист «Все свидетельства»", "все записи; данные карточки — где загружена"),
        ("Ограничение", "марки СМ извлечены автоматически из текста карточки; «и другие аналоги по ПТД» не раскрывается"),
    ]:
        ws.append([lab, val])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 30, 110
    for c in ws["A"]:
        c.font = Font(bold=True)
    _sheet(wb.create_sheet("Карточки"), H, [row(k) for k in ready], W, link_col=20)
    _sheet(wb.create_sheet("Марки СМ"), ["Марка СМ", "Свидетельств", "Организаций", "Группы ТУ", "Организации"],
           mk_rows, [30, 12, 12, 26, 100])
    _sheet(wb.create_sheet("Организации"), ["Организация", "Свидетельств ВП", "Карточек загружено",
                                           "Марки СМ (по загруженным)", "Способы", "Группы ТУ",
                                           "Основные материалы (примеры)"], org_rows, [45, 12, 12, 40, 24, 18, 80])
    _sheet(wb.create_sheet("Все свидетельства"), H,
           [row(k) for k in sorted(hs, key=lambda k: (hs[k]["organization"], k))], W, link_col=20)
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    wb.save(a.out)
    print(f"{a.out}: свидетельств {len(hs)}, организаций {len(by_org)}, карточек с данными {len(ready)}, "
          f"марок СМ {len(mk_rows)}")


if __name__ == "__main__":
    main()
